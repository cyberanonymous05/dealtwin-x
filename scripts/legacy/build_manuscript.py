"""Build the ACISM 2026 submission manuscript (.docx), matching the confirmed Springer LNCS
template specs: A4, single column, Times New Roman, 14pt bold title, 12pt bold L1 headings,
10pt bold L2 headings, 9pt abstract/keywords/captions/references, ~10pt body text.

Content reflects the project's ACTUAL, honestly-reported results -- nothing here overstates
what the experiments showed. See data/processed/ for the underlying JSON this is built from.
"""

import json
import os

from docx import Document
from docx.shared import Pt, Mm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

RESULTS_DIR = os.path.join("data", "processed")
# self-contained: writes inside this project regardless of where it's placed, so re-running after
# extracting/moving the project folder elsewhere still works without a sibling folder assumption
_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(_PROJECT_ROOT, "paper_assets", "paper", "ACISM2026_DealTwinX_Manuscript.docx")
FIGURES_DIR = os.path.join(_PROJECT_ROOT, "paper_assets", "figures")

FONT = "Times New Roman"


def load(name):
    with open(os.path.join(RESULTS_DIR, name)) as f:
        return json.load(f)


def set_margins_a4(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(1.73)
    section.right_margin = Inches(1.73)
    section.top_margin = Inches(2.05)
    section.bottom_margin = Inches(2.05)


def add_para(doc, text="", size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p


def add_mixed_para(doc, parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8):
    """parts: list of (text, bold, italic, size)"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold, italic, size in parts:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p


def add_heading1(doc, text):
    return add_para(doc, text, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def add_heading2(doc, text):
    return add_para(doc, text, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def add_figure(doc, path, caption, width_in=5.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    add_mixed_para(doc, [(caption[0], True, False, 9), (caption[1], False, False, 9)],
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(9)


def add_table(doc, headers, rows, caption):
    add_mixed_para(doc, [(caption[0], True, False, 9), (caption[1], False, False, 9)],
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.name = FONT
        run.font.size = Pt(9)
    for row_vals in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row_vals):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.name = FONT
            run.font.size = Pt(9)
    add_para(doc, "", size=4, space_after=10)


def main():
    pipeline = load("pipeline_results.json")
    ablation = load("ablation_and_baseline_results.json")
    learned = load("learned_metric_results.json")
    outcome = load("outcome_grounded_results.json")

    doc = Document()
    set_margins_a4(doc)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)

    # ---- Title / authors ----
    add_para(doc, "DealTwin-X: An Empirical Study of Multi-Modal Embeddings for M&A Precedent "
                  "Transaction Retrieval Against Naive Industry-Size Baselines",
             size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    add_para(doc, "[Author Name]¹", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "¹ [Author Affiliation, Institution, City, Country]", size=9,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "adeshkanna50@gmail.com", size=9, italic=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    # ---- Abstract ----
    add_mixed_para(doc, [
        ("Abstract. ", True, False, 9),
        ("Precedent transaction analysis -- estimating a target company's valuation multiple from "
         "comparable historical M&A deals -- is normally performed manually against proprietary "
         "databases such as Bloomberg and Capital IQ, filtered by industry and deal size alone. We "
         "present DealTwin-X, a pipeline built exclusively on free public data (SEC EDGAR full-text "
         "search, SEC XBRL Company Facts, FRED) that represents each deal as a multi-block \"DealDNA\" "
         "embedding -- industry, financial fundamentals, transaction structure, semantically-selected "
         "legal clauses (FinBERT), and macroeconomic context -- retrieved via FAISS, with an "
         "XGBoost/SHAP surrogate for the implied valuation multiple. Rather than reporting only "
         "favorable metrics, we evaluate DealDNA against a naive industry-and-size retriever using "
         "four complementary tests over a live corpus of 200 real merger filings (2016-2025, 44 "
         "industries): a heuristic relevance construct, a within-industry differentiation metric, a "
         "cross-validated learned block-weighting scheme (avoiding circularity), and a decisive test "
         "in which relevance is grounded in deals' REALIZED valuation multiples rather than any "
         "hand-built criterion. DealDNA is competitive with the naive baseline under the "
         "outcome-grounded test (winning or tying on 5 of 7 metrics), but learned block weights "
         "converge, under two independent relevance constructs, on essentially zero contribution from "
         "legal-clause embeddings. We report this negative finding transparently and discuss what "
         "would be required to overturn it.", False, False, 9),
    ], space_after=6)
    add_mixed_para(doc, [
        ("Keywords: ", True, False, 9),
        ("Precedent Transaction Analysis, Mergers and Acquisitions, Information Retrieval, FinBERT, "
         "Learning to Rank, Explainable AI, Financial NLP.", False, False, 9),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)

    # ---- 1. Introduction ----
    add_heading1(doc, "1   Introduction")
    add_para(doc, "Precedent transaction analysis is one of the three standard valuation methodologies "
                  "used in mergers and acquisitions (M&A) advisory, alongside discounted cash flow and "
                  "trading comparables analysis. An analyst identifies historical M&A deals judged "
                  "\"comparable\" to a target company, observes the valuation multiples those deals "
                  "commanded, and infers a justified multiple range for the current transaction. In "
                  "practice this search is performed against proprietary terminals (Bloomberg, "
                  "Capital IQ) using industry classification and deal-size filters, with little "
                  "systematic use of transaction structure, legal-clause content, or the macroeconomic "
                  "conditions prevailing at announcement.")
    add_para(doc, "This paper asks a narrower and more falsifiable question than \"can AI replace this "
                  "process\": does a multi-modal transaction embedding that incorporates financial "
                  "fundamentals, deal structure, legal-clause semantics, and macro context retrieve "
                  "better precedents than the naive industry-and-size filtering a banker already "
                  "performs for free? We built a complete, reproducible pipeline to answer it using "
                  "only free public data sources, and we report the answer honestly, including the "
                  "parts that did not support our initial hypothesis.")
    add_para(doc, "The contributions of this paper are: (1) a live ingestion pipeline over SEC EDGAR, "
                  "SEC XBRL, and FRED that discovers and enriches real merger filings without any "
                  "hand-curated deal list; (2) DealDNA, a five-block transaction embedding combining "
                  "industry identity, financial fundamentals, deal structure, semantically-selected "
                  "legal-clause passages embedded with FinBERT, and macroeconomic/temporal context; "
                  "(3) an XGBoost valuation-multiple surrogate explained with SHAP; (4) a four-stage "
                  "evaluation methodology -- naive-baseline comparison, within-industry differentiation, "
                  "cross-validated learned block weighting, and an outcome-grounded relevance test based "
                  "on realized valuation multiples -- designed specifically to avoid the circularity risk "
                  "of hand-built relevance labels; and (5) a transparent report of what the evidence does "
                  "and does not support, including a negative finding for the paper's original "
                  "legal-clause-embedding hypothesis.")
    add_para(doc, "Section 2 reviews related work in financial NLP and learning-to-rank. Section 3 "
                  "describes the data pipeline and DealDNA architecture. Section 4 presents the "
                  "experimental setup, hyperparameters, and results, including the ablation and "
                  "comparative analyses. Section 5 concludes with an honest discussion of what would be "
                  "needed to resolve the open questions this study raises.")

    # ---- 2. Literature Review ----
    add_heading1(doc, "2   Literature Review")
    add_para(doc, "Pre-trained language models transformed text representation learning following the "
                  "introduction of BERT [1]. Domain-adapted variants followed for finance specifically: "
                  "FinBERT models fine-tuned on financial communications and analyst-tone corpora [2] "
                  "provide sentence-level embeddings sensitive to financial terminology (e.g. "
                  "\"termination fee\", \"material adverse effect\") that general-domain models handle "
                  "less precisely. Sentence-level embedding models more broadly, such as Sentence-BERT "
                  "[3], enable efficient semantic similarity search over large text corpora via "
                  "pooled, fixed-size representations -- the approach we use to classify merger-filing "
                  "clauses by type without relying on literal keyword matches.")
    add_para(doc, "Efficient nearest-neighbor search over dense embeddings at the scale needed for a "
                  "growing transaction corpus is provided by FAISS [4], which we use as the retrieval "
                  "backbone. For the valuation-multiple surrogate and its explanation, we use XGBoost "
                  "[5], a gradient-boosted tree ensemble well suited to small, heterogeneous tabular "
                  "datasets, together with SHAP [6] for consistent, additive feature attribution.")
    add_para(doc, "Evaluating a ranked retrieval list against graded relevance is standard practice in "
                  "information retrieval via NDCG [7]; we adopt NDCG alongside MAP, MRR, precision, and "
                  "recall. Because no publicly available, expert-labeled relevance judgments exist for "
                  "M&A precedent comparability, we additionally draw on learning-to-rank methodology [8] "
                  "to fit retrieval weights directly from pairwise preference data, subject to "
                  "cross-validation, rather than relying solely on a hand-authored relevance function.")
    add_para(doc, "Comparable-company and precedent-transaction selection in traditional valuation "
                  "practice relies on standardized industry classification (SIC or GICS codes) and "
                  "size screens, refined manually by an analyst's judgment of qualitative similarity "
                  "in business model, growth profile, and deal structure [9]. This manual, judgment-"
                  "heavy step is precisely what a learned multi-modal embedding could in principle "
                  "systematize -- the premise this paper tests empirically rather than assumes.")
    add_para(doc, "To our knowledge, no prior published work evaluates a multi-modal transaction "
                  "embedding for M&A precedent retrieval against a naive industry-and-size baseline "
                  "using a corpus built entirely from free public filings; we position this paper's "
                  "contribution primarily as that empirical evaluation, rather than as a claim that the "
                  "embedding architecture itself is novel.")

    # ---- 3. Proposed Methodology ----
    add_heading1(doc, "3   Proposed Methodology")

    add_heading2(doc, "3.1   Data Set")
    add_para(doc, "The deal universe is discovered live from the SEC EDGAR full-text search API "
                  "(efts.sec.gov), querying for DEFM14A/DEFM14C merger proxy filings between 2016 and "
                  "2025 across five query phrasings designed to surface merger-agreement language "
                  "(e.g. \"merger agreement termination fee\", \"agreement and plan of merger material "
                  "adverse effect\"). Duplicate filers are removed and up to 200 unique deals are "
                  "retained. For each discovered deal we pull: (i) financial fundamentals (revenue, "
                  "EBITDA proxy, leverage, shares outstanding) from the SEC XBRL Company Facts API, "
                  "correctly restricted to annual-duration facts to avoid quarterly figures mistakenly "
                  "tagged under the same annual-report form; (ii) macroeconomic context (Fed funds "
                  "rate, 10-year Treasury yield, yield-curve slope, high-yield spread, VIX) from FRED "
                  "as of the filing date; and (iii) the full filing text itself. The resulting corpus "
                  "spans 200 real transactions across 44 distinct two-digit SIC industry groups, with "
                  "88% revenue coverage, 92% enterprise-value-estimate coverage, and 99% cash-per-share "
                  "coverage.")

    add_table(doc, ["Field", "Coverage"], [
        ["Revenue (annual, XBRL)", "88%"],
        ["EBITDA proxy (operating income + D&A)", "74%"],
        ["Revenue growth (YoY)", "86%"],
        ["Shares outstanding", "93%"],
        ["Enterprise value estimate", "92%"],
        ["Cash consideration per share", "99%"],
        ["Termination fee amount", "66%"],
        ["Stated deal premium (filing-disclosed)", "26%"],
    ], ("Table 0. ", "Field coverage across the 200-deal corpus. Fields not derivable for a given "
                     "deal (e.g. a stock-for-stock deal has no cash-per-share figure, and smaller "
                     "deals often omit a quantified fairness-opinion premium) are median-imputed "
                     "downstream rather than dropped, following Sect. 3.3."))

    add_heading2(doc, "3.2   Preprocessing")
    add_para(doc, "Filing HTML is converted to plain text and cleaned of markup and HTML entities. "
                  "Numeric deal terms -- cash consideration per share, stock exchange ratio, "
                  "termination fee amount, and consideration mix (cash / stock / mixed) -- are "
                  "extracted via targeted regular expressions tolerant of the line-wrapping common in "
                  "converted filing text. The deal premium is parsed directly from the fairness-opinion "
                  "disclosure every proxy contains (e.g. \"a 77% premium to the closing price ... the "
                  "last trading day prior to the announcement\"), tried in both common word orders, "
                  "which raised premium coverage from 3% (Yahoo Finance price history, unavailable for "
                  "most delisted targets) to 26%. Legal-clause passages for the embedding are NOT "
                  "selected by keyword matching. Instead, candidate sentences are segmented from the "
                  "filing, filtered to plausible legal prose, embedded with a sentence-transformer "
                  "model, and matched by cosine similarity against short prototype descriptions of "
                  "eight clause categories (material adverse effect, termination fee, earn-out / "
                  "contingent value right, no-solicitation / fiduciary-out, representations and "
                  "warranties, regulatory approval, appraisal rights, exchange-ratio consideration). "
                  "This recovers clauses phrased without the literal category keyword and correctly "
                  "separates distinct fee provisions (e.g. a financing-related termination fee from an "
                  "antitrust-related one) that a keyword window conflates.")

    add_heading2(doc, "3.3   Models")
    add_para(doc, "DealDNA represents each transaction as the concatenation of five blocks. The "
                  "industry block is a one-hot encoding of the two-digit SIC group. The financial "
                  "block (9 features) covers log revenue, revenue growth, EBITDA margin, net margin, "
                  "leverage, cash ratio, EV/Revenue, EV/EBITDA, and deal premium. The structure block "
                  "(6 features) covers consideration mix, earn-out presence, termination fee as a "
                  "percentage of enterprise value, and MAC-clause mention intensity. The macro block "
                  "(6 features) covers the five FRED series plus filing-date ordinal (added so the "
                  "embedding can directly capture deal recency, rather than only indirectly through "
                  "correlated rate levels). The legal block is the mean-pooled, 768-dimensional FinBERT "
                  "embedding of the selected clause passages. Numeric blocks are median-imputed and "
                  "z-scored corpus-wide; the legal block is L2-normalized and down-weighted so its 768 "
                  "dimensions do not dominate cosine similarity against the roughly 20 numeric "
                  "dimensions.")
    add_para(doc, "Retrieval is performed with a FAISS flat inner-product index over the (optionally "
                  "block-reweighted) embeddings. A gradient-boosted regressor (XGBoost) is trained to "
                  "predict log(EV/Revenue) from the non-multiple DealDNA features, explained "
                  "per-prediction and globally with SHAP. Because no expert relevance judgments are "
                  "available, block-level retrieval weights are additionally learned via a pairwise "
                  "ranking logistic regression: for sampled candidate pairs with differing relevance "
                  "under a given relevance construct, the model learns non-negative weights that best "
                  "separate more-relevant from less-relevant candidates, fit strictly on a GroupKFold "
                  "training split (grouped by two-digit SIC) and evaluated only on the held-out fold's "
                  "queries -- so the reported weights and metrics are not circular with each other.")

    # ---- 4. Experimental Results ----
    add_heading1(doc, "4   Experimental Results")

    add_heading2(doc, "4.1   Experimental Setup")
    add_para(doc, f"The final corpus contains {pipeline['n_deals']} deals; {pipeline['embedding_dim']}-"
                  "dimensional full DealDNA embeddings are built per deal. Retrieval quality is scored "
                  "under two independent relevance constructs. The heuristic construct grades relevance "
                  "in [0,3]: an industry match is a gate (0 if the two deals are in unrelated "
                  "industries), plus bonuses for similar revenue scale, matching consideration "
                  "structure, and filing recency -- modeled on how a banker actually screens "
                  "comparables. The outcome-grounded construct instead grades relevance by how similar "
                  "the two deals' REALIZED EV/Revenue multiples turned out to be (again industry-gated), "
                  "which is the actual premise of precedent transaction analysis and is not a criterion "
                  "either retriever is trivially built to satisfy. For this test the multiple itself is "
                  "excluded from DealDNA's features to avoid circularity. All cross-validated results use "
                  "5-fold GroupKFold splitting grouped by two-digit SIC, so an entire held-out industry's "
                  "relevance structure is never seen while fitting weights for it.")

    add_table(doc,
              ["Hyperparameter", "Value"],
              [
                  ["XGBoost: n_estimators / max_depth / learning_rate", "60 / 2 / 0.10"],
                  ["XGBoost: subsample / reg_lambda / min_child_weight", "0.9 / 2.0 / 2"],
                  ["Ridge baseline: alpha", "5.0"],
                  ["GroupKFold: n_splits (grouped by 2-digit SIC)", "5"],
                  ["Legal block weight (hand-picked, main pipeline)", "0.35"],
                  ["Industry block weight (hand-picked, main pipeline)", "1.5"],
                  ["Clause classifier: similarity threshold / top-n per type", "0.35 / 2"],
              ],
              ("Table 1. ", "Model and evaluation hyperparameters."))

    add_heading2(doc, "4.2   Results: Valuation-Multiple Surrogate")
    surr = pipeline["valuation_surrogate"]
    ridge = surr["model_comparison_ridge_baseline"]
    add_para(doc, f"The XGBoost surrogate, trained on {surr['n_training_deals']} deals with computable "
                  f"EV/Revenue, achieves a cross-validated R² of {surr['cv_mean_r2_log_space']:.2f} in log "
                  f"space (MAE {surr['cv_mean_mae_ev_to_revenue']:.2f}x on the raw multiple), against "
                  f"{ridge['cv_mean_r2_log_space']:.2f} R² for a Ridge baseline on identical folds and "
                  "features (Fig. 5) -- evidence the tree ensemble's added capacity is earning its keep "
                  "at this sample size rather than merely fitting noise. Fig. 4 shows the global SHAP "
                  "feature attribution: net margin, revenue scale, and EBITDA margin dominate, with "
                  "MAC-clause intensity and macro rate levels contributing modestly. The absolute MAE "
                  "remains large relative to typical multiples (median approximately 1.3x), so this "
                  "surrogate should be read as a directionally-informative explanation aid rather than a "
                  "precise valuation instrument at the current corpus size.")
    add_figure(doc, os.path.join(FIGURES_DIR, "fig4_shap_drivers.png"),
               ("Fig. 4. ", "Global SHAP feature importance for the XGBoost valuation-multiple surrogate."))
    add_figure(doc, os.path.join(FIGURES_DIR, "fig5_model_comparison.png"),
               ("Fig. 5. ", "Cross-validated MAE and R²: XGBoost vs. a Ridge baseline on identical folds."))

    add_heading2(doc, "4.3   Ablation Studies")
    add_para(doc, "Table 2 and Fig. 1 show within-industry NDCG@5 as DealDNA blocks are added "
                  "cumulatively, with hand-picked block weights, against the naive baseline. Financial "
                  "fundamentals alone already recover most of the achievable differentiation; industry "
                  "alone is the weakest single block for within-industry ranking (it cannot distinguish "
                  "between same-industry candidates at all, since a one-hot industry vector is "
                  "identical for every member of a group). Structure and macro contribute small, "
                  "consistent increments; adding the legal block produces no measurable further "
                  "improvement.")
    ndcg5 = "within_industry_ndcg@5"
    add_table(doc, ["Configuration", "Within-industry NDCG@5"], [
        ["Financial only", f"{ablation['financial_only'][ndcg5]:.3f}"],
        ["Industry only", f"{ablation['industry_only'][ndcg5]:.3f}"],
        ["Industry + Financial", f"{ablation['industry_plus_financial'][ndcg5]:.3f}"],
        ["+ Structure", f"{ablation['industry_financial_structure'][ndcg5]:.3f}"],
        ["+ Macro", f"{ablation['industry_financial_structure_macro'][ndcg5]:.3f}"],
        ["Full DealDNA (+ Legal)", f"{ablation['full_dealdna'][ndcg5]:.3f}"],
        ["Naive industry+size baseline", f"{ablation['naive_industry_size_baseline'][ndcg5]:.3f}"],
    ], ("Table 2. ", "Block ablation, hand-picked weights, heuristic relevance construct."))
    add_figure(doc, os.path.join(FIGURES_DIR, "fig1_ablation_progression.png"),
               ("Fig. 1. ", "Block ablation progression vs. the naive industry+size baseline."))

    add_para(doc, "Hand-picked block weights are an obvious source of bias, so block weights were "
                  "additionally learned via the cross-validated pairwise procedure of Sect. 3.3, under "
                  "both relevance constructs (Fig. 2). Under both constructs independently, the learned "
                  "weight for the legal block converges to approximately zero (0.006-0.008 out of a "
                  "5.0 total budget), while industry dominates the learned weighting overwhelmingly "
                  "(4.3-4.7 of that budget). This is the paper's central empirical finding: it is not a "
                  "hand-tuning artifact, since it emerges independently from two different, "
                  "non-circular, cross-validated fitting procedures.")
    add_figure(doc, os.path.join(FIGURES_DIR, "fig2_learned_weights.png"),
               ("Fig. 2. ", "Cross-validated learned block weights under two independent relevance constructs."))

    add_heading2(doc, "4.4   Comparative Analysis")
    add_para(doc, f"Under the heuristic relevance construct, the naive industry+size baseline scores "
                  f"{ablation['naive_industry_size_baseline'][ndcg5]:.3f} within-industry NDCG@5 against "
                  f"{ablation['full_dealdna'][ndcg5]:.3f} for full DealDNA with hand-picked weights and "
                  f"{learned['within_industry_ndcg@5']:.3f} with learned weights -- DealDNA does not "
                  "beat this baseline under this construct. However, this construct's relevance labels "
                  "directly mirror the two features the naive baseline itself encodes (industry, size), "
                  "which risks favoring the baseline by design rather than by genuine retrieval quality.")
    add_para(doc, "The outcome-grounded test (Table 3, Fig. 3), which grades relevance by realized "
                  "EV/Revenue multiple similarity rather than by industry/size directly, gives a "
                  "materially different picture: DealDNA (learned weights, multiples excluded from its "
                  "own features) wins or ties the naive baseline on 5 of 7 metrics, and is within 0.001 "
                  "on within-industry NDCG@10. This indicates the earlier result was, at least in part, "
                  "an artifact of how relevance was defined rather than proof DealDNA has no retrieval "
                  "value.")
    obase = outcome["naive_industry_size_baseline"]
    odna = outcome["full_dealdna_learned_weights_no_multiples"]
    add_table(doc, ["Metric", "Naive baseline", "DealDNA (learned, no multiples)"], [
        ["NDCG@5", f"{obase['ndcg@5']:.3f}", f"{odna['ndcg@5']:.3f}"],
        ["MAP", f"{obase['map']:.3f}", f"{odna['map']:.3f}"],
        ["MRR", f"{obase['mrr']:.3f}", f"{odna['mrr']:.3f}"],
        ["Precision@5", f"{obase['precision@5']:.3f}", f"{odna['precision@5']:.3f}"],
        ["Recall@10", f"{obase['recall@10']:.3f}", f"{odna['recall@10']:.3f}"],
        ["Within-industry NDCG@5", f"{obase['within_industry_ndcg@5']:.3f}", f"{odna['within_industry_ndcg@5']:.3f}"],
        ["Within-industry NDCG@10", f"{obase['within_industry_ndcg@10']:.3f}", f"{odna['within_industry_ndcg@10']:.4f}"],
    ], ("Table 3. ", "Outcome-grounded evaluation (realized-multiple relevance), n=135 scorable deals."))
    add_figure(doc, os.path.join(FIGURES_DIR, "fig3_outcome_grounded_comparison.png"),
               ("Fig. 3. ", "DealDNA vs. naive baseline under outcome-grounded (realized-multiple) relevance."))
    add_para(doc, "Reading both tests together: DealDNA's competitiveness is driven overwhelmingly by "
                  "industry identification and, to a much smaller extent, financial fundamentals and "
                  "macro timing -- not by the legal-clause embeddings that were this project's original, "
                  "most novel hypothesis. That hypothesis is not supported by either test conducted "
                  "here.")

    add_heading2(doc, "4.5   Threats to Validity")
    add_para(doc, "Internal validity. The two relevance constructs used here are both proxies for a "
                  "banker's actual comparability judgment, not that judgment itself; no free, "
                  "expert-labeled dataset of \"true\" precedent-transaction comparables exists to our "
                  "knowledge. The heuristic construct's overlap with the naive baseline's own features "
                  "(industry, size) is a known limitation we addressed directly with the "
                  "outcome-grounded construct, but that construct in turn depends on our own EV "
                  "estimation formula (cash-per-share x shares outstanding + debt - cash), which is a "
                  "simplification that can misstate enterprise value for asset-light or "
                  "financial-services targets holding large customer-settlement cash balances.")
    add_para(doc, "Construct validity. Numeric deal-term extraction (cash-per-share, exchange ratio, "
                  "termination fee, stated premium) is regex-based and can select the wrong figure when "
                  "a filing discloses multiple similar amounts, such as separate company-side and "
                  "parent-side termination fees. The semantic clause classifier, while shown "
                  "qualitatively to recover clauses a keyword search misses, has not itself been "
                  "validated against a labeled clause-type ground truth.")
    add_para(doc, "External validity. The corpus (200 deals, 44 industries, 2016-2025) is a live, "
                  "unbiased sample by discovery method, but is still small relative to the full "
                  "population of U.S. public-company mergers over that period, and is limited to "
                  "issuers required to file DEFM14A/DEFM14C proxies -- excluding, for instance, "
                  "all-cash tender offers without a proxy vote and non-U.S. targets. The "
                  "outcome-grounded test further restricts to the 135 deals with both a same-industry "
                  "peer and a computable multiple, a moderate sample for detecting a subtle, "
                  "second-order legal-clause effect if one exists at a smaller magnitude than industry "
                  "or size effects.")

    # ---- 5. Conclusion ----
    add_heading1(doc, "5   Conclusion")
    add_para(doc, "We built and evaluated DealTwin-X, a free-data pipeline for M&A precedent "
                  "transaction retrieval, using a live corpus of 200 real SEC merger filings and a "
                  "four-stage evaluation methodology designed specifically to avoid the circularity "
                  "risks of a single hand-built relevance function. Under an outcome-grounded test in "
                  "which relevance reflects deals' actual realized valuation multiples, DealDNA is "
                  "competitive with a naive industry-and-size baseline, winning or tying on 5 of 7 "
                  "retrieval metrics -- a genuine, if modest, positive result. Under both relevance "
                  "constructs tested, however, cross-validated learned block weights place essentially "
                  "no value on the legal-clause FinBERT embeddings that motivated this project's "
                  "original framing, while industry identity dominates and financial fundamentals "
                  "contribute a smaller, consistent amount.")
    add_para(doc, "We treat this as an honest, reportable negative finding rather than as a reason to "
                  "keep searching for a more favorable evaluation. Two limitations bound how strongly it "
                  "should be read: first, no expert-labeled relevance judgments exist for M&A "
                  "comparability from free sources, so both relevance constructs used here remain "
                  "defensible proxies rather than ground truth; second, the outcome-grounded test scores "
                  "only 135 of 200 deals (those with both a same-industry peer and a computable "
                  "multiple), a moderate sample for detecting a subtler, second-order legal-clause "
                  "effect. Future work should prioritize obtaining even a small set of practitioner-"
                  "adjudicated relevance judgments that explicitly credit legal and structural "
                  "similarity -- the one signal type neither test used here -- alongside scaling the "
                  "corpus further, before drawing a final conclusion about whether legal-clause "
                  "embeddings add value to precedent transaction retrieval.")

    # ---- References ----
    add_heading1(doc, "References")
    refs = [
        "Devlin, J., Chang, M.W., Lee, K., Toutanova, K.: BERT: Pre-training of Deep Bidirectional "
        "Transformers for Language Understanding. In: Proceedings of NAACL-HLT 2019, pp. 4171-4186 (2019).",
        "Yang, Y., Uy, M.C.S., Huang, A.: FinBERT: A Pretrained Language Model for Financial "
        "Communications. arXiv preprint arXiv:2006.08097 (2020).",
        "Reimers, N., Gurevych, I.: Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. In: "
        "Proceedings of EMNLP-IJCNLP 2019, pp. 3982-3992 (2019).",
        "Johnson, J., Douze, M., Jégou, H.: Billion-scale similarity search with GPUs. IEEE "
        "Transactions on Big Data 7(3), 535-547 (2019).",
        "Chen, T., Guestrin, C.: XGBoost: A Scalable Tree Boosting System. In: Proceedings of the 22nd "
        "ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, pp. 785-794 (2016).",
        "Lundberg, S.M., Lee, S.I.: A Unified Approach to Interpreting Model Predictions. In: Advances "
        "in Neural Information Processing Systems 30 (NeurIPS 2017), pp. 4765-4774 (2017).",
        "Järvelin, K., Kekäläinen, J.: Cumulated Gain-Based Evaluation of IR Techniques. "
        "ACM Transactions on Information Systems 20(4), 422-446 (2002).",
        "Liu, T.Y.: Learning to Rank for Information Retrieval. Foundations and Trends in Information "
        "Retrieval 3(3), 225-331 (2009).",
        "Damodaran, A.: Investment Valuation: Tools and Techniques for Determining the Value of Any "
        "Asset. 3rd edn. Wiley, Hoboken, NJ (2012).",
        "U.S. Securities and Exchange Commission: EDGAR Full-Text Search and XBRL Company Facts APIs, "
        "https://www.sec.gov/edgar, last accessed 2026/08/21.",
        "Federal Reserve Bank of St. Louis: FRED Economic Data, https://fred.stlouisfed.org, last "
        "accessed 2026/08/21.",
    ]
    for i, ref in enumerate(refs, 1):
        add_mixed_para(doc, [(f"[{i}] ", False, False, 9), (ref, False, False, 9)],
                        align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    doc.save(OUT_PATH)
    print(f"Saved manuscript to {OUT_PATH}")


if __name__ == "__main__":
    main()
